# -*- coding: utf-8 -*-
import os, zipfile, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

D = r"L:\giehub-test\分析\Parameter-Plan-V1000"
OUT = os.path.join(D, "_extract")
os.makedirs(OUT, exist_ok=True)

files = [
    "V1000完整论文.docx",
    "V1000 系统(已有版权禁止商用).docx",
    "v1000流程.docx",
    "v1000示范.docx",
    "技术论文",
]

import docx

for f in files:
    p = os.path.join(D, f)
    base = os.path.splitext(f)[0]
    if not os.path.exists(p):
        print("MISSING", f)
        continue
    print("=" * 60)
    print("FILE:", f, os.path.getsize(p))
    if not zipfile.is_zipfile(p):
        with open(p, encoding="utf-8", errors="replace") as fh:
            txt = fh.read()
        out = os.path.join(OUT, base + ".txt")
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(txt)
        print("plain text ->", out, len(txt), "chars")
        continue
    doc = docx.Document(p)
    lines = []
    for para in doc.paragraphs:
        style = (para.style.name or "") if para.style is not None else ""
        t = para.text.replace("\xa0", " ")
        if not t.strip():
            continue
        if style.startswith("Heading") or style.startswith("标题"):
            lines.append("\n## [%s] %s" % (style, t.strip()))
        else:
            lines.append(t)
    for ti, tbl in enumerate(doc.tables):
        lines.append("\n[TABLE %d]" % ti)
        for row in tbl.rows:
            lines.append(" | ".join(c.text.replace("\n", " ").replace("\xa0", " ") for c in row.cells))
    # headers/footers? skip
    txt = "\n".join(lines)
    out = os.path.join(OUT, base + ".txt")
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(txt)
    print("docx ->", out, len(txt), "chars")
    print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))

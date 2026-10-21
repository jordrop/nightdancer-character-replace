# -*- coding: utf-8 -*-
"""提取 docx 结构：标题层级 + 段落数 + 章节前若干字，便于选段。"""
import os
import sys
import io

from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

FILES = [
    (r"E:\伊雪的活\女红军.docx", "女红军"),
    (r"E:\伊雪的活\新刻拍案惊奇（博士国）书稿.docx", "博士国"),
]
OUT = r"L:\giehub-test\分析\归墟选段"
os.makedirs(OUT, exist_ok=True)

for path, tag in FILES:
    doc = Document(path)
    paras = doc.paragraphs
    print("=" * 70)
    print(tag, "paragraphs:", len(paras), "tables:", len(doc.tables))
    total_chars = sum(len(p.text) for p in paras)
    print("总字符:", total_chars)
    # 保存全文（带样式前缀，供后续检索）
    lines = []
    headings = []
    for p in paras:
        t = (p.text or "").replace("\xa0", " ")
        if not t.strip():
            continue
        st = p.style.name if p.style is not None else ""
        if st.startswith("Heading") or st.startswith("标题") or (len(t) < 60 and (
                t.startswith("第") and ("章" in t or "回" in t or "节" in t)) or t.startswith("序") or t.startswith("尾声")):
            headings.append((st, t))
            lines.append("\n## [%s] %s" % (st, t))
        else:
            lines.append(t)
    full = os.path.join(OUT, tag + "_全文.txt")
    with open(full, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("全文 ->", full, len("\n".join(lines)), "chars")
    print("--- 识别到的标题/章节（前 60） ---")
    for st, t in headings[:60]:
        print("[%s] %s" % (st, t))
    print("--- 标题总数:", len(headings))
